from pathlib import Path
from docx import Document
import openpyxl
root=Path(__file__).resolve().parents[2]
for p in [root/'ruanjianshixun/需求分析/02_语法分析模块需求分析.docx',root/'ruanjianshixun/分模块问答整理/02_语法分析.docx']:
    print('\nSOURCE',p.name)
    d=Document(p)
    for para in d.paragraphs:
        if para.text.strip(): print(para.text)
    for t in d.tables:
        for row in t.rows: print(' / '.join(c.text for c in row.cells))
for p in (root/'canKao').glob('*.xlsx'):
    w=openpyxl.load_workbook(p,data_only=True)
    print('\nREFERENCE',p.name,[(s.title,s.max_row,s.max_column) for s in w])
    for s in w.worksheets[:1]:
        for row in s.iter_rows(min_row=1,max_row=min(12,s.max_row),values_only=True):print(row[:6])
