from pathlib import Path
import json,copy
from openpyxl import load_workbook
root=Path(__file__).parent
before=load_workbook(root/'scope_edit_before.xlsx',data_only=False)
after=load_workbook(root.parent/'LL1与图示/C89_LL1集合与预测分析表.xlsx',data_only=False)
edits={(s,a):v for s,a,v in json.loads((root/'scope_edits.json').read_text('utf-8'))}
assert before.sheetnames==after.sheetnames
unchanged=0;changed=0
for old,new in zip(before,after):
    assert old.max_row==new.max_row and old.max_column==new.max_column
    assert old.freeze_panes==new.freeze_panes
    assert list(old.tables)==list(new.tables)
    for table in old.tables:assert old.tables[table].ref==new.tables[table].ref
    for row in old:
        for cell in row:
            other=new[cell.coordinate];key=(old.title,cell.coordinate)
            if key in edits:assert other.value==edits[key];changed+=cell.value!=other.value
            else:assert cell.value==other.value,(key,cell.value,other.value);unchanged+=1
            assert copy.copy(cell.font)==copy.copy(other.font),(key,'font')
            assert copy.copy(cell.fill)==copy.copy(other.fill),(key,'fill')
            assert copy.copy(cell.alignment)==copy.copy(other.alignment),(key,'alignment')
            assert cell.number_format==other.number_format,(key,'number_format')
report={'changed_text_cells':changed,'unchanged_cells_checked':unchanged,'grammar_and_sets_preserved':True,'cell_formatting_preserved':True,'tables_and_freeze_panes_preserved':True}
(root/'scope_validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),'utf-8')
print(json.dumps(report,ensure_ascii=False))
