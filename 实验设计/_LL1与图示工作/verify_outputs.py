import json
from pathlib import Path
from openpyxl import load_workbook

root=Path(__file__).parent
core=json.loads((root/'core_grammar_data.json').read_text('utf-8'))
raw=json.loads((root/'grammar_data.json').read_text('utf-8'))
loc=json.loads((root/'sheet_locations.json').read_text('utf-8'))
wb=load_workbook(root.parent/'LL1与图示/C89_LL1集合与预测分析表.xlsx',data_only=True)
checks=0
def eq(a,b):
    global checks
    assert a==b,(a,b)
    checks+=1
def rows(name):
    return list(wb[name].iter_rows(min_row=loc[name]['header']+1,max_row=loc[name]['last'],values_only=True))
def fmt(s): return '{ '+', '.join(s)+' }'
for n,r in zip(core['grammar'],rows('FIRST集')):
    eq(r[0],n);eq(r[1],fmt(core['first'][n]))
for n,r in zip(core['grammar'],rows('FOLLOW集')):
    eq(r[0],n);eq(r[1],fmt(core['follow'][n]))
for p,r in zip(core['productions'],rows('SELECT集')):
    eq(list(r[:5]),[p['id'],p['lhs'],' '.join(p['rhs']) or 'ε',fmt(p['first_rhs']),fmt(p['select'])])
for p,r in zip(core['productions'],rows('核心文法')):
    eq(list(r[:3]),[p['id'],p['lhs'],' '.join(p['rhs']) or 'ε'])
tm={(t['lhs'],t['lookahead']):' / '.join(t['productions']) for t in core['table']}
for n,r in zip(core['grammar'],rows('预测分析表')):
    eq(r[0],n)
    for t,v in zip(core['terminals'],r[1:]):eq(v,tm.get((n,t),'—'))
for n,r in zip(raw['grammar'],rows('C89原始集合')):
    eq(list(r[:3]),[n,fmt(raw['first'][n]),fmt(raw['follow'][n])])
for p,r in zip(raw['productions'],rows('C89原始SELECT')):
    eq(list(r[:5]),[p['id'],p['lhs'],' '.join(p['rhs']) or 'ε',fmt(p['first_rhs']),fmt(p['select'])])
for c,r in zip(raw['conflicts'],rows('冲突分析')):
    eq(list(r[:3]),[c['lhs'],c['lookahead'],' / '.join(c['productions'])])
dfa=json.loads((root.parent/'_状态转换表工作/dfa_data.json').read_text('utf-8'))
literal={r['token']:r.get('literal') or r['name'] for r in dfa['rules'] if r.get('token')}
literal.update({t:w for w,t in dfa['keywords'].items()})
for t,r in zip(core['terminals'],rows('终结符映射')):
    eq(r[0],t)
    if t in literal:eq(r[1],literal[t])
errors=[(s.title,c.coordinate,c.value) for s in wb for row in s for c in row if c.data_type=='e']
eq(errors,[])
report={'saved_xlsx_comparisons':checks,'formula_errors':errors,'sheets':wb.sheetnames,'core_conflicts':len(core['conflicts']),'raw_conflicts':len(raw['conflicts'])}
(root/'xlsx_validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),'utf-8')
print(json.dumps(report,ensure_ascii=False))
print('DFA keys',list(dfa['ordinary']))
for i,(a,w) in enumerate(zip(dfa['ordinary']['accept'],dfa['ordinary']['witnesses'])):
    rule=dfa['rules'][a] if a>=0 else {}
    print(i,repr(w),rule.get('name'),rule.get('token'))
print('header keys',list(dfa['header']))
print('rule sample',dfa['rules'][:3])
